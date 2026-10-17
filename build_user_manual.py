from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path('THN_Attendance_User_Manual.docx')
LOGO = Path('frontend/employee_grouping_system/static/images/logo.png')

NAVY = '0B2545'; BLUE = '2E74B5'; DARKBLUE = '1F4D78'; LIGHT = 'E8EEF5'; PALE = 'F4F6F9'; GRAY = '5B6573'; GOLD = '7A5A00'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None: node = OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None: tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    tblInd = tblPr.first_child_found_in('w:tblInd')
    if tblInd is None: tblInd = OxmlElement('w:tblInd'); tblPr.append(tblInd)
    tblInd.set(qn('w:w'), str(indent)); tblInd.set(qn('w:type'),'dxa')
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths): col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcW = cell._tc.tcPr.tcW; tcW.set(qn('w:w'), str(width)); tcW.set(qn('w:type'),'dxa')
            set_cell_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); item = OxmlElement('w:tblHeader'); item.set(qn('w:val'),'true'); trPr.append(item)

def set_run(run, size=11, color=NAVY, bold=None, italic=None):
    run.font.name = 'Calibri'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def set_para(p, before=0, after=6, line=1.25, align=None):
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = line
    if align is not None: p.alignment = align

def style_doc(doc):
    sec = doc.sections[0]; sec.top_margin=Inches(0.8); sec.bottom_margin=Inches(0.72); sec.left_margin=Inches(0.85); sec.right_margin=Inches(0.85); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
    normal = doc.styles['Normal']; normal.font.name='Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(NAVY); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
    for name, size, color, before, after in [('Title',28,NAVY,0,8),('Subtitle',13,GRAY,0,16),('Heading 1',16,BLUE,18,8),('Heading 2',13,BLUE,12,6),('Heading 3',11.5,DARKBLUE,8,4)]:
        s=doc.styles[name]; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=(name!='Subtitle'); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
    for n in ['List Bullet','List Number']:
        s=doc.styles[n]; s.font.name='Calibri'; s.font.size=Pt(10.5); s.paragraph_format.space_after=Pt(3); s.paragraph_format.line_spacing=1.2
    for sname in ['Note Box','Small']:
        if sname not in [s.name for s in doc.styles]: doc.styles.add_style(sname, WD_STYLE_TYPE.PARAGRAPH)
    note=doc.styles['Note Box']; note.font.name='Calibri'; note.font.size=Pt(10); note.font.color.rgb=RGBColor.from_string(DARKBLUE); note.paragraph_format.space_before=Pt(5); note.paragraph_format.space_after=Pt(6); note.paragraph_format.line_spacing=1.15
    small=doc.styles['Small']; small.font.name='Calibri'; small.font.size=Pt(8.5); small.font.color.rgb=RGBColor.from_string(GRAY); small.paragraph_format.space_after=Pt(3)

def add_header_footer(section):
    h=section.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_para(h,0,0,1); r=h.add_run('TIME HUB NETWORK  |  ATTENDANCE MONITORING SYSTEM'); set_run(r,8,GRAY,True)
    f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(f,0,0,1); r=f.add_run('Internal user guide  •  Version 1.0'); set_run(r,8,GRAY)

def add_title(doc, text, subtitle=None):
    p=doc.add_paragraph(style='Heading 1'); p.add_run(text)
    if subtitle:
        p=doc.add_paragraph(); set_para(p,0,9,1.15); r=p.add_run(subtitle); set_run(r,10,GRAY,italic=True)

def add_bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); set_para(p,0,3,1.2); p.add_run(item)

def add_steps(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Number'); set_para(p,0,3,1.2); p.add_run(item)

def note(doc, title, text):
    t=doc.add_table(rows=1, cols=1); set_table_geometry(t,[9360]); t.alignment=WD_TABLE_ALIGNMENT.LEFT; set_cell_shading(t.cell(0,0), PALE)
    set_repeat_table_header(t.rows[0])
    p=t.cell(0,0).paragraphs[0]; p.style='Note Box'; r=p.add_run(title+'  '); set_run(r,10,DARKBLUE,True); r=p.add_run(text); set_run(r,10,DARKBLUE)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def simple_table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.style='Table Grid'; set_table_geometry(t,widths)
    for c, text in zip(t.rows[0].cells,headers):
        set_cell_shading(c,LIGHT); p=c.paragraphs[0]; set_para(p,0,0,1.1); r=p.add_run(text); set_run(r,9.5,DARKBLUE,True)
    set_repeat_table_header(t.rows[0])
    for row in rows:
        cells=t.add_row().cells
        for c, text in zip(cells,row):
            p=c.paragraphs[0]; set_para(p,0,0,1.15); r=p.add_run(text); set_run(r,9.5,NAVY)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

doc=Document(); style_doc(doc); add_header_footer(doc.sections[0])

# Cover
doc.add_paragraph().paragraph_format.space_after=Pt(40)
if LOGO.exists():
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(LOGO), width=Inches(0.9)); doc.inline_shapes[-1]._inline.docPr.set('descr', 'Time Hub Network logo')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(p,18,8,1); r=p.add_run('TIME HUB NETWORK'); set_run(r,12,BLUE,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(p,0,8,1); r=p.add_run('Attendance Monitoring System'); set_run(r,28,NAVY,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(p,0,30,1.2); r=p.add_run('User Manual Guide'); set_run(r,16,GRAY)
simple_table(doc,['Guide purpose','Audience'],[['How to navigate, monitor attendance, manage workforce records, and use approval workflows.','Authorized users: Viewers, Supervisors, HR/Administration, and System Administrators.']],[4680,4680])
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(p,16,0,1); r=p.add_run('Version 1.0  |  Internal use'); set_run(r,9,GRAY,italic=True)
doc.add_page_break()

add_title(doc,'1. Getting started','What you need before you begin')
doc.add_paragraph('The Attendance Monitoring System is a web-based workforce console for reviewing attendance, organizing personnel, maintaining employee records, and processing attendance corrections. Your available buttons and records depend on your account role, department, and assigned scope.')
doc.add_heading('Sign in and sign out', level=2)
add_steps(doc,['Open the system URL provided by your administrator and enter your username and password.','After successful sign-in, use the left navigation panel to open the required module. Your name and assigned department appear at the top of the panel.','When finished, select Logout at the bottom of the navigation panel. Do not leave an active session open on a shared computer.'])
note(doc,'Access-aware navigation.','If a button or menu item is not visible, it is likely outside your assigned role or department scope. Contact an Administrator rather than using another person’s account.')
doc.add_heading('Navigation at a glance', level=2)
simple_table(doc,['Module','Primary use'],[['Dashboard','Review attendance rates and weekly/monthly trend summaries.'],['Grouping','Organize staff by line, department, and shift; import or move manpower when authorized.'],['Employee Management','Find, review, create, update, or maintain employee records.'],['Attendance Logs','Review daily records, filter attendance, submit corrections, and complete approvals.'],['Scanner Management','Admin-only configuration of biometric scanner inventory.']],[2640,6720])

add_title(doc,'2. Roles and access scope','Use the right workflow for your authorization')
simple_table(doc,['User type','Typical capability'],[['Viewer','Read-only access to authorized records; can review employee details and attendance history.'],['Supervisor','Operational workforce tasks within assigned scope, including approved imports and transfers.'],['HR / Administration','May complete final attendance-correction approval and perform administration-authorized work.'],['Administrator','Full system administration, including user provisioning, employee maintenance, and scanner setup.']],[2640,6720])
doc.add_paragraph('Access may also be limited by department, production line, shift, or production section. For example, a line-level user sees only the assigned line and shift, while a section-level Production user sees only the assigned section.')
note(doc,'Important.','Permissions are enforced by the system. Do not attempt to work around a missing action; send the request to the correct approver or system administrator.')

add_title(doc,'3. Dashboard: monitor attendance','Use summaries to spot trends and exceptions')
doc.add_paragraph('Open Dashboard from the Overview menu to review attendance-rate summaries. The dashboard provides broad weekly and monthly views, including breakdowns that may be grouped by agency, tenure, department, or line.')
add_bullets(doc,['Start with the current period and identify unusually low attendance rates or high late/absent counts.','Use line or department breakdowns to focus follow-up on the affected operational area.','Treat dashboard figures as monitoring indicators. Use Attendance Logs to review the underlying employee records and punch details.','Where configured, review the daily validation indicator for THN Day Shift; it compares the actual present/late/absent total with the expected benchmark headcount.'])
doc.add_heading('Recommended daily routine', level=2)
add_steps(doc,['Review the dashboard for exceptions at the beginning of the shift.','Open Attendance Logs and filter to the relevant date, shift, line, or status.','Investigate missing punches and pending correction requests.','Escalate unresolved discrepancies through the approved attendance-correction process.'])

add_title(doc,'4. Workforce grouping','Review lines, shifts, and manpower assignments')
doc.add_paragraph('The Grouping workspace organizes employees into line and shift groups. It can show departmental groupings, line summaries, unassigned employees, and movement history. The available editing controls are restricted to authorized users.')
doc.add_heading('Find and review a group', level=2)
add_steps(doc,['Open Grouping from the Workforce menu.','Select the applicable shift and use the search or line filters to narrow the view.','Expand a line group to see its employees, or select Summary to review line-level totals.','Select an employee name to review recorded movement/history information.'])
doc.add_heading('Manage manpower (authorized users)', level=2)
add_steps(doc,['Select the employees to move.','Choose Manage Manpower or the available transfer control.','Choose the destination sheet or line and section. If permitted, enter a new section name.','Confirm the transfer and then verify the employee appears in the correct destination group.'])
note(doc,'Audit trail.','Line, shift, and process changes detected by imports or updates are retained as movement/history records, helping managers trace workforce changes.')
doc.add_heading('Import an Excel staffing file', level=2)
add_steps(doc,['In Grouping, select Import Excel.','Choose the approved workbook and select whether to apply it immediately or schedule it for a future date and time.','Confirm the import result and review any skipped sheets or error messages.','Validate the affected lines and employee movements after the import is active.'])
doc.add_paragraph('For best results, staffing sheets should be named with DS or NS (Day Shift or Night Shift) and include an Emp No. header. Scheduled imports activate automatically at the chosen date and time.')

add_title(doc,'5. Employee management','Maintain complete and accurate employee records')
doc.add_paragraph('Employee Management is the system directory for active workforce records. Use filters and search to locate a worker, click a row to view details, and use editing controls only when your role allows changes.')
doc.add_heading('Add a new employee', level=2)
add_steps(doc,['Open Employee Management and select Add New Employee.','Enter the required identity and employment details, such as name, agency, employee number, department, position, hire date, and shift.','Set the production line where applicable. This field is generally used for non-Administration departments.','Save the record and confirm that the employee appears in the directory and appropriate grouping.'])
doc.add_paragraph('Tenure is calculated automatically from the hire date; do not enter it manually.')
doc.add_heading('Edit, leave, and deactivate records', level=2)
add_bullets(doc,['Select Edit on an employee row to update employee data. Changes are recorded in the audit trail with old and new values.','To mark staff On Leave in bulk, select the employees, provide a reason, and optionally provide a return date. Use Return to Work to clear the leave status when they resume duty.','Bulk Delete sets an employee to INACTIVE rather than permanently erasing the record, preserving historical data. This is an Administrator-only action.'])
note(doc,'Data-quality check.','Before saving a change, confirm the employee number, agency, department, line, shift, and hire date. These fields affect grouping, reporting, and attendance interpretation.')

add_title(doc,'6. Attendance logs and corrections','Review records, resolve exceptions, and approve changes')
doc.add_paragraph('Attendance Logs contains daily time-in/time-out records and calculated attendance status. Use this module for detailed follow-up, not just summary reporting.')
doc.add_heading('Review and filter attendance', level=2)
add_steps(doc,['Open Attendance Logs from the Time Tracking menu.','Set filters for shift, date, status, position, or use free-text search by employee ID or name. Combine filters to focus on a specific group.','Select a column header to sort results.','Select an employee name or row to open available attendance history and statistics.'])
simple_table(doc,['Status','Meaning / action'],[['Present','Expected attendance condition is met.'],['Late','Arrival is past the shift cutoff; confirm whether a correction is required.'],['Absent','No qualifying attendance record for the shift; verify with the responsible supervisor.'],['On Leave','Employee is marked on leave in the directory; verify return-to-work status when appropriate.'],['Pending','A relevant attendance action or correction is still awaiting completion.']],[2520,6840])
doc.add_heading('Handle a missing punch or correction', level=2)
add_steps(doc,['Use the Missing Punches report for the selected date and shift to identify records with only time-in or time-out recorded.','Open the employee’s attendance record and select Edit or the correction action, if enabled for your role.','Enter the correct information and submit the request with accurate supporting details.','Monitor the approval queue or notifications until the request is resolved.'])
note(doc,'Two-stage approval.','Attendance corrections first go to the applicable department, line, or section approver. HR/Final Approval then completes the change; the live attendance record is updated only after final approval. Viewers cannot submit edits or act on approval requests.')

add_title(doc,'7. Scanner management (Administrators)','Maintain the biometric device inventory')
doc.add_paragraph('Scanner Management is visible only to Administrators and requires an additional scanner-access passcode. Use it to maintain the registered biometric device inventory used by attendance synchronization.')
add_steps(doc,['Open Scanner Management from System Setup and complete the additional access prompt.','Review the registered scanner list, including agency, IP address, scanner type, location, port, and enabled status.','Add or edit a scanner using a valid IP address, supported type (ZKTeco or Anviz), existing agency, optional port, and location.','Disable a scanner when it should stop syncing but its historical records must remain available.'])
note(doc,'Preserve history.','A scanner that already has biometric logs cannot be deleted. Disable it instead, so attendance history remains traceable.')

add_title(doc,'8. Troubleshooting and good practice','Keep records dependable and secure')
simple_table(doc,['Situation','What to do'],[['A button or page is unavailable','Confirm your role and assigned scope. Ask an Administrator to review access if the action is required for your job.'],['Excel import skips a sheet','Confirm the sheet name contains DS or NS and its header includes Emp No.; correct the file and import again.'],['A staff member is missing from a line','Search Employee Management, confirm active status and shift, then use authorized grouping/transfer actions.'],['A punch is missing or wrong','Use Attendance Logs, submit a correction request, and follow it through both approval stages.'],['Scanner cannot be removed','Disable it when it contains biometric history; only unused scanners may be deleted.'],['Attendance count does not look right','Check filters, date and shift; compare the dashboard indicator with detailed Attendance Logs before escalating.']],[3200,6160])
doc.add_heading('Operational checklist', level=2)
add_bullets(doc,['Use the correct shift and date before taking action.','Check employee identity and agency before changing a record.','Use imports and bulk actions only with an approved source file and verified scope.','Review pending corrections and missing punches routinely.','Log out of shared devices and never share credentials or scanner-access passcodes.'])
doc.add_heading('Need help?', level=2)
doc.add_paragraph('Use the in-system assistant for common how-to questions about imports, transfers, access levels, employee records, and attendance approvals. For permissions, data corrections outside your scope, or technical issues, contact your designated Administrator or Supervisor.')

doc.add_page_break()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(p,120,10,1); r=p.add_run('End of guide'); set_run(r,16,BLUE,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_para(p,0,0,1.2); r=p.add_run('Time Hub Network Attendance Monitoring System'); set_run(r,10,GRAY)
doc.save(OUT)
print(OUT.resolve())
