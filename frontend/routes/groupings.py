from flask import Blueprint, request, jsonify, session
from sqlalchemy import text as sa_text
from io import BytesIO
from dateutil import parser as dateparser
from openpyxl import load_workbook
import traceback
import re
import json
from datetime import datetime, timezone, timedelta
import os
import secrets
import tempfile
import time
from routes.extensions import db
from routes.decorators import login_required, admin_required, roles_required
from routes.auth import (manage_employees_required, can_view_employee,
                         can_manage_employees_check, parse_employee_department)
from routes.emp_id import build_emp_id, agency_code_from_emp_id, AGENCY_ID as AGENCY_ID_MAP

groupings_bp = Blueprint("groupings", __name__)

# Same lock as attendance_processor: history and the derived attendance must
# not be changed concurrently by an import/manual move and a processor run.
ATTENDANCE_LOCK = 4815162342
PLANT_TZ = timezone(timedelta(hours=8))


def plant_now():
    return datetime.now(PLANT_TZ).replace(tzinfo=None)


def parse_schedule_effective_at(raw, mode, effective_at):
    if raw:
        parsed = datetime.strptime(raw, "%Y-%m-%d")
        if parsed.strftime("%Y-%m-%d") != raw:
            raise ValueError("Use YYYY-MM-DD")
        return parsed
    if mode == "immediate":
        return effective_at.replace(hour=0, minute=0, second=0, microsecond=0)
    return effective_at  # Preserve explicit scheduled-time behavior.


def queue_attendance_rebuild(conn, employee_ids, effective_at):
    """Requeue raw history from the neighboring overnight shift forward.

    No attendance status is overwritten here. The processor consults the full
    dated schedule and protects manual_override records. Shared transaction.
    """
    employee_ids = sorted(set(employee_ids))
    if not employee_ids:
        return 0
    result = conn.execute(sa_text("""
        UPDATE backend.biometric_logs SET processed = FALSE
        WHERE emp_id = ANY(:ids) AND punch_time >= :since
    """), {"ids": employee_ids,
           "since": datetime.combine(effective_at.date(), datetime.min.time()) - timedelta(days=1)})
    return result.rowcount


def queue_batch_attendance(conn, batch_id, employees):
    effective_at = conn.execute(sa_text("""
        SELECT COALESCE(schedule_effective_at, effective_at) FROM frontend.import_batches WHERE id = :id
    """), {"id": batch_id}).scalar()
    if effective_at is None:
        raise ValueError("Batch effective_at is required for attendance history")
    return queue_attendance_rebuild(conn,
        [str(e.get("employee_id", "")).strip().upper() for e in employees if e.get("employee_id")],
        effective_at)


# ─── Constants ───────────────────────────────────────────────────
COLUMN_ALIASES = {
    "hc": "hc",

    "last name": "last_name",
    "lastname": "last_name",
    "lname": "last_name",

    "first name": "first_name",
    "firstname": "first_name",
    "fname": "first_name",

    "id no": "employee_id",
    "id no.": "employee_id",
    "id number": "employee_id",
    "emp no": "employee_id",
    "emp no.": "employee_id",
    "employee no": "employee_id",
    "employee number": "employee_id",
    "empid": "employee_id",

    "full name": "full_name",
    "fullname": "full_name",
    "name": "full_name",

    "date hired": "date_hired",

    "process": "process",
    "line": "line",

    "emp": "emp",
    "emp.": "emp",
    "employment": "emp",

    "tenure": "tenure",

    "department": "department",
    "dept": "department",

    "position": "position",

    "machine": "machine",
    
    "wing": "wing",

    "remarks": "remarks",
    "remark": "remarks",

    "gender": "gender",
    "sex": "gender",
}

KNOWN_AGENCIES = {"THN", "CCI", "NATCORP", "EMIS"}

TRACKED_FIELDS = {
    "shift":   "SHIFT",
    "process": "PROCESS",
    "line":    "LINE",
}

def normalize(value: object) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    return "" if s.upper() in ("N/A", "NA", "-", "NONE", "NULL") else s


def clean(value: object) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return None if s.upper() in ("", "N/A", "NA", "-", "NONE") else s


def safe_date(value: object):
    if not value:
        return None
    try:
        return dateparser.parse(str(value)).date()
    except Exception:
        return None


# Maps the sheet's short M/F convention to the same Male/Female values
# Employee Management's add/edit dropdowns store, so both entry paths
# produce identical-looking values in the gender column instead of
# silently forking into two spellings for the same thing. Anything
# already spelled out (or unrecognized) passes through normalize()
# unchanged rather than being dropped.
GENDER_ALIASES = {"M": "Male", "F": "Female"}

def normalize_gender(value: object) -> str:
    s = normalize(value)
    return GENDER_ALIASES.get(s.upper(), s)


# ─── Excel helpers ────────────────────────────────────────────────
def parse_sheet_name(sheet_name: str) -> dict:
    raw = sheet_name.strip().upper()

    if re.search(r'\bNS\b', raw):
        return {
            "line_id": raw,
            "shift": "Night Shift"
        }

    if re.search(r'\bDS\b', raw):
        return {
            "line_id": raw,
            "shift": "Day Shift"
        }

    raise ValueError(
        f"Sheet '{sheet_name}' skipped — no DS or NS found."
    )


def extract_sub_tables(ws) -> list:
    # read_only worksheets trust the <dimension> tag in the file, which some
    # exporters write wrong (e.g. A1:A1). Reset so every row is read.
    if hasattr(ws, "reset_dimensions"):
        ws.reset_dimensions()
    all_rows = list(ws.iter_rows(values_only=True))

    header_row = None
    header_index = None

    for i, row in enumerate(all_rows):
        cells = [
            str(c).strip().lower()
            for c in row
            if c is not None and str(c).strip()
        ]

        if "emp no." in cells or "emp" in cells:
            header_row = list(row)
            header_index = i
            break

    if not header_row:
        return []

    data_rows = []

    for row in all_rows[header_index + 1:]:
        if not row:
            continue

        if all(c is None or str(c).strip() == "" for c in row):
            continue

        data_rows.append(list(row))


    return [{
        "agency": None,
        "header_row": header_row,
        "data_rows": data_rows
    }]

def normalize_id(value: object) -> str:
    """
    Excel numeric ID cells are read by openpyxl as float (1222 -> 1222.0),
    so a naive str() on the cell value stores "1222.0" while every other
    system (attendance-logs API, employee master) uses "1222". Strip a
    trailing ".0" off a purely-numeric string so IDs match across sources.
    """
    s = str(value or "").strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s

def parse_block(block: dict, line_id: str, shift: str, sheet_name: str) -> tuple:
    header   = block.get("header_row", [])
    agency   = block.get("agency", "UNKNOWN")
    warnings: list[str] = []

    if not header:
        warnings.append(f"Sheet '{sheet_name}' → Agency '{agency}': empty header, skipped.")
        return [], warnings

    col_map: dict[str, int] = {}

    for col_idx, cell in enumerate(header):
            if cell is None:
                continue

            key = str(cell).strip().lower()

            if key in COLUMN_ALIASES:
                col_map[COLUMN_ALIASES[key]] = col_idx

    if "employee_id" not in col_map:
        for i, cell in enumerate(header):
            if cell is None:
                continue
            val = str(cell).strip().lower()
            if "last" in val:                                          col_map.setdefault("last_name",   i)
            elif "first" in val:                                       col_map.setdefault("first_name",  i)
            elif "emp no" in val or "emp. no" in val or "no." in val: col_map.setdefault("employee_id", i)
            elif "full" in val:                                        col_map.setdefault("full_name",   i)
            elif "date" in val or "hired" in val:                      col_map.setdefault("date_hired",  i)
            elif "process" in val:                                     col_map.setdefault("process",     i)
            elif "remarks" in val:                                     col_map.setdefault("remarks",     i)
            elif "tenure" in val or val in ("ten.", "ten"):            col_map.setdefault("tenure",      i)
            elif "line" in val:                                        col_map.setdefault("line",        i)
            elif val in ("emp", "emp."):                               col_map.setdefault("emp",         i)
            elif val in ("hc", "h.c", "head"):                        col_map.setdefault("hc",          i)
            elif "machine" in val:                                    col_map.setdefault("machine",     i)
            elif val == "wing":                                       col_map.setdefault("wing",        i)
            elif "gender" in val or val == "sex":                      col_map.setdefault("gender",      i)

    if "employee_id" not in col_map:
        for field, pos in [
            ("hc", 0), ("last_name", 1), ("first_name", 2), ("employee_id", 3),
            ("full_name", 4), ("date_hired", 5), ("process", 6),
            ("line", 7), ("emp", 8), ("tenure", 9),
        ]:
            col_map.setdefault(field, pos)

    employees: list[dict] = []

    for raw_row in block.get("data_rows", []):
        if not raw_row:
            continue

        def get(field: str, row: list = raw_row) -> object:
            idx = col_map.get(field)
            return None if idx is None or idx >= len(row) else row[idx]

        def get_text(field: str) -> str:
            val = get(field)
            return "" if val is None else str(val).strip()

        emp_id_raw = normalize_id(get_text("employee_id"))
        if not emp_id_raw or emp_id_raw.lower() in ("none", "nan", "-", "n/a", "emp no.", "emp no"):
            continue

        agency_code = get_text("emp").upper().strip()
        emp_id      = build_emp_id(agency_code, emp_id_raw)
        last        = get_text("last_name")
        first       = get_text("first_name")
        full        = get_text("full_name") or f"{first} {last}".strip()

        tenure_raw = get("tenure")
        tenure: float | None = None
        if tenure_raw not in ("", None):
            if isinstance(tenure_raw, (int, float)):
                tenure = float(tenure_raw)
            else:
                try:
                    tenure = float(str(tenure_raw).strip())
                except ValueError:
                    tenure = None

        employees.append({
            "employee_id": emp_id.strip().upper(),
            "emp_no":      emp_id_raw,             # raw scanner ID, before agency prefix
            "last_name":   last,   "first_name": first,
            "full_name":   full,   "name":       full,
            "agency":      agency_code,
            "section":     get_text("process") or get_text("section"),
            "sheet_name":  sheet_name,
            "line_id":     str(line_id).strip(),
            "shift":       str(shift).strip(),
            "hc":          get_text("hc"),
            "date_hired":  get_text("date_hired"),
            "process":     normalize(get_text("process")),
            "department": normalize(get_text("department")),
            "position": normalize(get_text("position")),
            "machine":  normalize(get_text("machine")),
            "wing":     normalize(get_text("wing")),
            "remarks": normalize(get_text("remarks")),
            "gender": normalize_gender(get_text("gender")),
            "line":        str(line_id).strip(),
            "sub_line":    normalize(get_text("line")),
            "emp":         get_text("emp"),
            "tenure":      tenure,

        })

    return employees, warnings


def detect_movements(conn, batch_id: int, all_employees: list) -> tuple:

    prev_batch = conn.execute(sa_text("""
        SELECT id FROM frontend.import_batches
        WHERE id < :current_id ORDER BY id DESC LIMIT 1
    """), {"current_id": batch_id}).fetchone()

    if not prev_batch:
        return 0, []

    prev_batch_id = prev_batch[0]
    prev_rows     = conn.execute(sa_text("""
        SELECT emp_id, full_name, shift, process, line
        FROM frontend.employee_snapshot_records
        WHERE batch_id = :batch_id
    """), {"batch_id": prev_batch_id}).fetchall()

    prev_map = {
        normalize(r[0]).upper(): {
            "full_name": r[1] or "",
            "shift":     normalize(r[2]),
            "process":   normalize(r[3]),
            "line":      normalize(r[4]),
        }
        for r in prev_rows
    }

    movement_count = 0
    shift_changed: list[dict] = []
    movement_rows: list[dict] = []

    for emp in all_employees:
        emp_id    = normalize(emp.get("employee_id", "")).upper()
        prev_data = prev_map.get(emp_id)
        if not emp_id or not prev_data:
            continue

        for field, field_name in TRACKED_FIELDS.items():
            curr_val = normalize(emp.get(field, ""))
            prev_val = prev_data.get(field, "")
            if not curr_val or not prev_val or curr_val.upper() == prev_val.upper():
                continue

            movement_rows.append({
                "emp_id": emp_id, "full_name": emp.get("full_name", ""),
                "field_name": field_name, "old_value": prev_val, "new_value": curr_val,
                "from_batch_id": prev_batch_id, "to_batch_id": batch_id,
            })
            movement_count += 1

            if field == "shift":
                shift_changed.append({
                    "emp_id":    emp_id,
                    "old_shift": prev_val,
                    "new_shift": curr_val,
                })

    # One round trip for all movements (was one INSERT per changed field).
    if movement_rows:
        conn.execute(sa_text("""
            INSERT INTO frontend.employee_movements
                (emp_id, full_name, field_name, old_value, new_value,
                 from_batch_id, to_batch_id, detected_at)
            SELECT x.emp_id, x.full_name, x.field_name, x.old_value, x.new_value,
                   x.from_batch_id, x.to_batch_id, NOW()
            FROM jsonb_to_recordset(CAST(:payload AS jsonb)) AS x(
                emp_id text, full_name text, field_name text,
                old_value text, new_value text,
                from_batch_id integer, to_batch_id integer)
        """), {"payload": json.dumps(movement_rows)})

    return movement_count, shift_changed


def fetch_current_state(conn) -> dict:
    """
    Builds the complete, global set of line/shift groups from the LIVE
    employee_groupings table — not the batch that was just imported.

    Because every employee row is keyed by emp_id and carries its own
    line_name/line/shift, grouping this table at query time is what makes
    the merge automatic:
      - Sheets from a previous import that weren't touched by this import
        simply remain in the result set (their employee rows still exist).
      - A sheet name that reappears in a new import isn't duplicated,
        because each employee within it was just UPSERTed via
        ON CONFLICT (emp_id) DO UPDATE — so the group it's collected into
        here reflects the new data in place.

    This replaces "return the batch" with "return the world" as the
    contract of the import endpoint.
    """
    rows = conn.execute(sa_text("""
        SELECT emp_id, emp_no, full_name, line_name, line, shift, process, position, machine, tenure, date_hired
        FROM frontend.employee_groupings
        WHERE COALESCE(employee_status, 'ACTIVE') <> 'INACTIVE'
        ORDER BY line_name, shift, last_name, first_name
    """)).mappings().all()

    groups_by_key: dict[str, dict] = {}
    unassigned: list[dict] = []

    for r in rows:
        emp_payload = {
            "id":      r["emp_id"],
            "name":    r["full_name"],
            "section": r["process"] or "",
            "process": r["process"] or "",
            "position": r["position"] or "",
            "machine": r["machine"] or "",
        }

        line_name = clean(r["line_name"])
        shift     = clean(r["shift"])

        if not line_name or not shift:
            unassigned.append(emp_payload)
            continue

        key = f"{line_name}||{shift}"
        if key not in groups_by_key:
            groups_by_key[key] = {
                "id":        key,          # stable across imports — doubles as merge key
                "line_name": line_name,
                "line":      r["line"],
                "shift":     shift,
                "employees": [],
            }
        groups_by_key[key]["employees"].append(emp_payload)

    groups = list(groups_by_key.values())
    for g in groups:
        g["count"] = len(g["employees"])

    return {"groups": groups, "unassigned": unassigned}
class _StepTimer:
    """Prints how long each import step took to the server console, so a slow
    import can be traced to the exact step (look for lines tagged [import-timing])."""
    def __init__(self, label: str):
        self.label = label
        self.t0 = self.last = time.perf_counter()
        self.steps: list = []

    def lap(self, name: str) -> None:
        now = time.perf_counter()
        self.steps.append((name, now - self.last))
        self.last = now

    def report(self, status: str = "ok") -> None:
        total = time.perf_counter() - self.t0
        print(f"[import-timing] {self.label} {status} total={total:.1f}s | "
              + " | ".join(f"{n}={d:.1f}s" for n, d in self.steps), flush=True)


# ─── Dry-run preview cache ───────────────────────────────────────
# The dry run already parses the whole workbook. Instead of making the browser
# upload + the server re-parse the same file for the real import, the parsed
# result is stashed on disk (shared by all workers on this host) under a
# one-time token. If the token is missing/expired/not yours, the route answers
# code="preview_expired" and the client simply re-sends the file.
PREVIEW_DIR = os.path.join(tempfile.gettempdir(), "import_previews")
PREVIEW_TTL_SECONDS = 15 * 60
_PREVIEW_TOKEN_RE = re.compile(r"^[A-Za-z0-9_-]{20,64}$")


def _purge_old_previews() -> None:
    now = time.time()
    try:
        names = os.listdir(PREVIEW_DIR)
    except OSError:
        return
    for name in names:
        path = os.path.join(PREVIEW_DIR, name)
        try:
            if now - os.path.getmtime(path) > PREVIEW_TTL_SECONDS:
                os.remove(path)
        except OSError:
            pass


def save_preview(payload: dict) -> str | None:
    try:
        os.makedirs(PREVIEW_DIR, mode=0o700, exist_ok=True)
        _purge_old_previews()
        token = secrets.token_urlsafe(24)
        tmp = os.path.join(PREVIEW_DIR, token + ".tmp")
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(payload, f, default=str)
        os.replace(tmp, os.path.join(PREVIEW_DIR, token + ".json"))
        return token
    except Exception:
        traceback.print_exc()
        return None  # caller just omits the token; client falls back to re-upload


def load_preview(token: str, username: str, agency_scope: str) -> dict | None:
    if not token or not _PREVIEW_TOKEN_RE.match(token):
        return None
    path = os.path.join(PREVIEW_DIR, token + ".json")
    try:
        if time.time() - os.path.getmtime(path) > PREVIEW_TTL_SECONDS:
            return None
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
    except (OSError, ValueError):
        return None
    # Bound to the user and scope the dry run was made with.
    if data.get("username") != username or data.get("agency_scope") != agency_scope:
        return None
    return data


def discard_preview(token: str) -> None:
    if token and _PREVIEW_TOKEN_RE.match(token):
        try:
            os.remove(os.path.join(PREVIEW_DIR, token + ".json"))
        except OSError:
            pass


# ─── Routes ──────────────────────────────────────────────────────
@groupings_bp.route("/api/import-grouping", methods=["POST"])
@roles_required("Admin", "Supervisor")
def import_grouping():
    try:
        # ── NEW: guard — roles_required already requires login, so
        # session["username"] should always be set here. If it's ever
        # missing, that's an auth/session bug — don't silently fall back
        # to a generic "system" label that would make the import log
        # untraceable to a real account.
        uploader_username = session.get("username")
        if not uploader_username:
            return jsonify({
                "success": False,
                "message": "Session error: uploader identity missing. Please log in again.",
            }), 401

        # A real import may reference the dry run's cached parse instead of
        # re-uploading the workbook (see "Dry-run preview cache" above).
        preview_token = (request.form.get("preview_token") or "").strip()
        use_preview = bool(preview_token) and request.form.get("dry_run") != "1"
        file = None
        source_filename = None
        if not use_preview:
            if "file" not in request.files:
                return jsonify({"success": False, "message": "No file uploaded."}), 400

            file = request.files["file"]
            if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
                return jsonify({"success": False, "message": "Please upload a valid Excel file."}), 400
            source_filename = file.filename

        # ── NEW: read effectivity choice from the form ──
        effective_mode = request.form.get("effective_mode", "immediate")  # "immediate" | "scheduled"
        effective_at_raw = request.form.get("effective_at")  # ISO string, only when scheduled

        if effective_mode not in ("immediate", "scheduled"):
            return jsonify({"success": False, "message": "Invalid effective mode."}), 400
        if effective_mode == "scheduled":
            if not effective_at_raw:
                return jsonify({"success": False, "message": "Please choose an effective date/time."}), 400
            try:
                effective_at = datetime.fromisoformat(effective_at_raw)
                if effective_at.tzinfo is not None:
                    effective_at = effective_at.astimezone(PLANT_TZ).replace(tzinfo=None)
            except ValueError:
                return jsonify({"success": False, "message": "Invalid effective date/time format."}), 400
        else:
            effective_at = plant_now()

        # Schedule validity is separate from the time the live roster changes.
        # Immediate uploads default to today's shifts, including morning arrivals.
        schedule_date_raw = request.form.get("schedule_effective_date", "").strip()
        try:
            schedule_effective_at = parse_schedule_effective_at(schedule_date_raw, effective_mode, effective_at)
        except ValueError:
            return jsonify({"success": False, "message": "Schedule effective date must be YYYY-MM-DD."}), 400

        # ── NEW: read agency scope choice from the form ──
        agency_scope = request.form.get("agency_scope", "ALL").strip().upper()
        if agency_scope != "ALL" and agency_scope not in KNOWN_AGENCIES:
            return jsonify({
                "success": False,
                "message": f"Invalid agency scope: {agency_scope}",
            }), 400

        # dry_run=1: preview lang kung sino ang matatanggal, walang isinusulat.
        # deactivate_missing=1: gawing INACTIVE ang mga wala na sa Excel.
        dry_run = request.form.get("dry_run") == "1"
        # The success payload used to always carry the whole live roster
        # (fetch_current_state). The web UI reloads it itself, so it can opt out.
        include_state = request.form.get("include_state", "1") != "0"
        # Level 1 (LINE) accounts ay hindi pwedeng mag-deactivate.
        can_deactivate = can_manage_employees_check()
        deactivate_missing = request.form.get("deactivate_missing") == "1" and can_deactivate
        uploader_access = _session_user_snapshot()

        result_lines  = []
        all_warnings  = []
        sheet_errors  = []
        all_employees = []
        # Mga sheet na may problema sa pag-parse: ang mga empleyado ng line na
        # ito ay hindi ide-deactivate (baka hindi lang nabasa ang sheet).
        problem_sheets: set = set()

        cached = None
        if use_preview:
            cached = load_preview(preview_token, uploader_username, agency_scope)
            if cached is None:
                return jsonify({
                    "success": False, "code": "preview_expired",
                    "message": "The checked file is no longer available. Please upload it again.",
                }), 410
            source_filename = cached.get("filename")

        if cached is not None:
            result_lines   = cached["result_lines"]
            all_warnings   = cached["warnings"]
            sheet_errors   = cached["errors"]
            all_employees  = cached["employees"]
            problem_sheets = set(cached["problem_sheets"])
        else:
            # read_only streams the XML instead of building every cell object —
            # much faster and lighter on memory for big workbooks.
            wb = load_workbook(BytesIO(file.read()), read_only=True, data_only=True)
            try:
                for sheet_name in wb.sheetnames:
                    try:
                        meta = parse_sheet_name(sheet_name)
                    except ValueError as e:
                        sheet_errors.append(str(e))
                        problem_sheets.add(sheet_name)
                        continue

                    try:
                        blocks = extract_sub_tables(wb[sheet_name])
                    except Exception as e:
                        sheet_errors.append(f"Sheet '{sheet_name}': {str(e)}")
                        problem_sheets.add(sheet_name)
                        continue

                    if not blocks:
                        all_warnings.append(f"Sheet '{sheet_name}': no data found.")
                        problem_sheets.add(sheet_name)
                        continue

                    sheet_employees: list = []
                    for block in blocks:
                        if not block.get("header_row"):
                            problem_sheets.add(sheet_name)
                        try:
                            emps, warns = parse_block(block, meta["line_id"], meta["shift"], sheet_name)
                            sheet_employees.extend(emps)
                            all_warnings.extend(warns)
                        except Exception as e:
                            all_warnings.append(f"Block error in '{sheet_name}': {str(e)}")
                            problem_sheets.add(sheet_name)

                    # ── NEW: keep only the rows matching the chosen agency scope ──
                    if agency_scope != "ALL":
                        sheet_employees = [
                            e for e in sheet_employees
                            if e.get("agency", "").upper() == agency_scope
                        ]

                    all_employees.extend(sheet_employees)
                    result_lines.append({
                        "line_name":      f"{meta['line_id']} — {meta['shift']}",
                        "line_id":        meta["line_id"],
                        "shift":          meta["shift"],
                        "sheet_name":     sheet_name,
                        "employee_count": len(sheet_employees),
                        "employees":      sheet_employees,
                    })

            finally:
                wb.close()

        if not result_lines and sheet_errors:
            return jsonify({
                "success": False,
                "message": "No valid sheets found.\n\n" + "\n".join(sheet_errors),
            }), 400

        # A single employee cannot have two assignments in the same batch.
        seen_assignments = {}
        for employee in all_employees:
            key = employee["employee_id"]
            assignment = (employee.get("shift"), employee.get("line"))
            if key in seen_assignments and seen_assignments[key] != assignment:
                return jsonify({"success": False, "message":
                    f"Conflicting assignments for {key}; correct the Excel file."}), 400
            seen_assignments[key] = assignment

        protected_sheets = sorted(problem_sheets)

        if dry_run:
            try:
                with db.engine.connect() as conn:
                    to_remove = find_employees_missing_from_import(
                        conn, all_employees, agency_scope, protected_sheets,
                        user=uploader_access) if can_deactivate else []
            except Exception as e:
                traceback.print_exc()
                return jsonify({"success": False, "message": f"Database error: {str(e)}"}), 500
            preview_token_out = save_preview({
                "username": uploader_username, "agency_scope": agency_scope,
                "filename": file.filename if file else source_filename,
                "result_lines": result_lines, "warnings": all_warnings,
                "errors": sheet_errors, "employees": all_employees,
                "problem_sheets": protected_sheets,
            })
            return jsonify({
                "success": True,
                "dry_run": True,
                "preview_token": preview_token_out,
                "agency_scope": agency_scope,
                "total_in_file": len(seen_assignments),
                "will_deactivate": to_remove,
                "can_deactivate": can_deactivate,
                "departments_in_file": sorted(departments_in_import(all_employees)),
                "protected_sheets": protected_sheets,
                "warnings": all_warnings, "errors": sheet_errors,
            })

        movements_count = 0
        shift_changed: list[dict] = []
        deactivated: list[dict] = []
        batch_id = None

        timer = _StepTimer(f"batch of {len(all_employees)} employees")
        try:
            with db.engine.begin() as conn:
                # Fail with a clear error instead of hanging for minutes when
                # another process (e.g. attendance processing) holds the lock.
                conn.execute(sa_text("SET LOCAL lock_timeout = '60s'"))
                timer.lap("parse+validate")
                conn.execute(sa_text("SELECT pg_advisory_xact_lock(:key)"), {"key": ATTENDANCE_LOCK})
                timer.lap("wait for attendance lock")
                ensure_batch_columns(conn)
                timer.lap("ensure columns")
                # ── NEW: batch insert now carries effective_at / applied / raw_payload ──
                is_immediate = (effective_mode != "scheduled")
                batch = conn.execute(sa_text("""
                    INSERT INTO frontend.import_batches
                        (filename, imported_by, total_lines, total_employees,
                         effective_at, schedule_effective_at, applied, raw_payload,
                         deactivate_scope, protected_sheets, deactivate_user)
                    VALUES (:filename, :imported_by, :total_lines, :total_employees,
                            :effective_at, :schedule_effective_at, :applied, :raw_payload,
                            :deactivate_scope, :protected_sheets, :deactivate_user)
                    RETURNING id
                """), {
                    "filename":        source_filename,
                    "imported_by":     uploader_username,   # ← real logged-in user, not "system"
                    "total_lines":     len(result_lines),
                    "total_employees": len(all_employees),
                    "effective_at":    effective_at,
                    "schedule_effective_at": schedule_effective_at,
                    "applied":         is_immediate,
                    "raw_payload":     json.dumps(all_employees),
                    "deactivate_scope": agency_scope if deactivate_missing else None,
                    "protected_sheets": json.dumps(protected_sheets),
                    "deactivate_user":  json.dumps(uploader_access) if deactivate_missing else None,
                })
                batch_id = batch.fetchone()[0]
                timer.lap("insert batch")

                # employee_snapshot_records is history regardless of mode — keep writing it.
                # One lookup + one INSERT for the whole batch (was 2 queries per employee).
                snapshot_ids = [str(e.get("employee_id", "")).strip().upper() for e in all_employees]
                snapshot_ids = [i for i in snapshot_ids if i]
                prev_by_id = {}
                if snapshot_ids:
                    prev_by_id = {
                        r[0]: r for r in conn.execute(sa_text(
                            "SELECT emp_id, process, line, shift "
                            "FROM frontend.employee_groupings WHERE emp_id = ANY(:ids)"
                        ), {"ids": snapshot_ids})
                    }

                snapshot_rows: dict[str, dict] = {}  # keyed by emp_id: last row wins, like the old per-row upsert
                for emp in all_employees:
                    emp_id        = str(emp.get("employee_id", "")).strip().upper()
                    if not emp_id:
                        continue
                    date_hired    = safe_date(emp.get("date_hired"))
                    process_value = clean(emp.get("process"))
                    line_value    = clean(emp.get("line"))
                    shift_value   = clean(emp.get("shift"))
                    emp_str       = str(emp.get("emp", "")).strip().upper()

                    prev = prev_by_id.get(emp_id)
                    if prev:
                        process_value = process_value or prev[1]
                        line_value    = line_value    or prev[2]
                        shift_value   = shift_value   or prev[3]

                    snapshot_rows[emp_id] = {
                        "batch_id": batch_id, "emp_id": emp_id,
                        "emp":       AGENCY_ID_MAP.get(emp_str),
                        "full_name": emp.get("full_name", ""),
                        "line":      line_value,    "shift":     shift_value,
                        "process":   process_value, "tenure":    emp.get("tenure"),
                        "date_hired": date_hired.isoformat() if date_hired else None,
                    }

                if snapshot_rows:
                    conn.execute(sa_text("""
                        INSERT INTO frontend.employee_snapshot_records
                            (batch_id, emp_id, emp, full_name, line, shift, process, tenure, date_hired)
                        SELECT x.batch_id, x.emp_id, x.emp, x.full_name, x.line, x.shift,
                               x.process, x.tenure, x.date_hired
                        FROM jsonb_to_recordset(CAST(:payload AS jsonb)) AS x(
                            batch_id integer, emp_id text, emp integer, full_name text,
                            line text, shift text, process text,
                            tenure double precision, date_hired date)
                        ON CONFLICT (batch_id, emp_id) DO UPDATE SET
                            full_name = EXCLUDED.full_name, line      = EXCLUDED.line,
                            shift     = EXCLUDED.shift,     process   = EXCLUDED.process,
                            tenure    = EXCLUDED.tenure,    date_hired= EXCLUDED.date_hired
                    """), {"payload": json.dumps(list(snapshot_rows.values()))})

                timer.lap("snapshot records")

                # ── NEW: only touch the live roster if this import is immediate ──
                if is_immediate:
                    _upsert_batch_to_roster(conn, batch_id, all_employees)
                    timer.lap("roster upsert")
                    deactivated = deactivate_missing_employees(conn, batch_id, all_employees)
                    timer.lap("deactivate missing")
                    movements_count, shift_changed = detect_movements(conn, batch_id, all_employees)
                    timer.lap("detect movements")
                    rows_recalculated = 0  # queued below; no current-shift historical overwrite
                else:
                    rows_recalculated = 0

                punches_queued = queue_batch_attendance(conn, batch_id, all_employees)
                timer.lap("queue attendance rebuild")
            timer.lap("commit")
            timer.report()

        except Exception as e:
            timer.report("FAILED")
            traceback.print_exc()
            return jsonify({"success": False, "message": f"Database error: {str(e)}"}), 500

        if cached is not None:
            discard_preview(preview_token)

        # Read the live roster after the import transaction (and its advisory
        # lock) has been released, and only if the caller wants it.
        full_state = {"groups": [], "unassigned": []}
        if include_state:
            try:
                with db.engine.connect() as conn:
                    full_state = fetch_current_state(conn)
            except Exception:
                traceback.print_exc()

        return jsonify({
            "success": True,
            "lines":        full_state["groups"],
            "unassigned":   full_state["unassigned"],
            "imported_sheets": result_lines,
            "batch_id": batch_id,
            "warnings": all_warnings, "errors": sheet_errors,
            "movements_found": movements_count,
            "shift_changes_recalculated": rows_recalculated,
            "attendance_punches_queued": punches_queued,
            "effective_mode": effective_mode,
            "effective_at": effective_at.isoformat(),
            "schedule_effective_at": schedule_effective_at.isoformat(),
            "agency_scope": agency_scope,
            "imported_by": uploader_username,   # ── NEW: surfaced to the frontend success dialog
            "deactivate_requested": deactivate_missing,
            "deactivated": deactivated,
            "deactivated_count": len(deactivated),
        })

    except Exception as e:
        traceback.print_exc()
        return jsonify({"success": False, "message": f"Server error: {str(e)}"}), 500
def _employee_rows(conn, batch_id: int, all_employees: list) -> list[dict]:
    """Normalize an import and fetch existing values with one round trip."""
    employee_ids = [str(emp.get("employee_id", "")).strip().upper() for emp in all_employees]
    employee_ids = [emp_id for emp_id in employee_ids if emp_id]
    existing = {}
    if employee_ids:
        rows = conn.execute(sa_text("""
            SELECT emp_id, process, line, shift, gender
            FROM frontend.employee_groupings
            WHERE emp_id = ANY(:employee_ids)
        """), {"employee_ids": employee_ids}).mappings()
        existing = {row["emp_id"]: row for row in rows}

    prepared = []
    for emp in all_employees:
        emp_id = str(emp.get("employee_id", "")).strip().upper()
        if not emp_id:
            continue
        prior = existing.get(emp_id, {})
        agency = str(emp.get("emp", "")).strip().upper()
        prepared.append({
            "batch_id": batch_id, "emp_id": emp_id, "emp": AGENCY_ID_MAP.get(agency),
            "emp_no": emp.get("emp_no", ""), "last_name": emp.get("last_name", ""),
            "first_name": emp.get("first_name", ""), "full_name": emp.get("full_name", ""),
            "line_name": emp.get("sheet_name", ""),
            "line": clean(emp.get("line")) or prior.get("line"),
            "shift": clean(emp.get("shift")) or prior.get("shift"),
            "process": clean(emp.get("process")) or prior.get("process"),
            "department": emp.get("department", ""), "position": emp.get("position", ""),
            "machine": emp.get("machine", ""), "wing": emp.get("wing", ""),
            "remarks": emp.get("remarks", ""), "tenure": emp.get("tenure"),
            "gender": clean(emp.get("gender")) or prior.get("gender"),
            "date_hired": safe_date(emp.get("date_hired")),
        })
    return prepared


# ── Agency → direct THN transfer handling ────────────────────────────────────
# Kapag ang agency worker (hal. "N12") ay naging direct THN (bagong ID, hal.
# "4587"), ang upsert ay ON CONFLICT (emp_id), kaya magkaibang row ang
# mabubuo para sa iisang tao at maiiwan ang luma. Ang lumang agency row ay
# minamarkahang INACTIVE (hindi binubura, para buo pa rin ang history).
# True = tumugma lang kapag pareho rin ang line_name (mas ligtas laban sa
# magkaibang taong magkapangalan, pero maaaring may matirang doble kapag
# lumipat din ng line ang empleyado).
TRANSFER_DEDUPE_REQUIRE_SAME_LINE = False


def _person_name_key(name: object) -> str:
    # "DELA CRUZ, JUAN" == "Juan Dela Cruz"
    tokens = re.findall(r"[A-Z0-9Ñ]+", str(name or "").upper())
    return " ".join(sorted(tokens))


def _find_transferred_stale(direct_rows, candidate_rows, exclude_ids=frozenset()) -> list[dict]:
    """Mga agency row (candidate_rows) na may THN row (direct_rows) na
    kapareho ng pangalan."""
    direct_by_name: dict[str, list] = {}
    for d in direct_rows:
        key = _person_name_key(d["full_name"])
        if key:
            direct_by_name.setdefault(key, []).append(d)

    stale = []
    for c in candidate_rows:
        cid = str(c["emp_id"] or "").strip().upper()
        if not cid or cid in exclude_ids or agency_code_from_emp_id(cid) == "THN":
            continue
        matches = direct_by_name.get(_person_name_key(c["full_name"]))
        if not matches:
            continue
        if TRANSFER_DEDUPE_REQUIRE_SAME_LINE:
            c_line = str(c["line_name"] or "").strip().upper()
            matches = [m for m in matches if str(m["line_name"] or "").strip().upper() == c_line]
            if not matches:
                continue
        stale.append({
            "emp_id": cid,
            "full_name": c["full_name"],
            "replaced_by": str(matches[0]["emp_id"]).strip().upper(),
        })
    return stale


def _mark_inactive(conn, emp_ids: list) -> None:
    if not emp_ids:
        return
    conn.execute(sa_text("""
        UPDATE frontend.employee_groupings
        SET employee_status = 'INACTIVE', updated_at = NOW()
        WHERE emp_id = ANY(:ids)
          AND COALESCE(employee_status, 'ACTIVE') = 'ACTIVE'
    """), {"ids": emp_ids})


def deactivate_transferred_agency_rows(conn, incoming_rows: list) -> list[dict]:
    """Tinatawag pagkatapos ng roster upsert. Kapag may THN row sa import na
    kapareho ng pangalan ng dating agency row, i-INACTIVE ang luma."""
    direct_rows = [r for r in incoming_rows
                   if agency_code_from_emp_id(str(r["emp_id"]).strip().upper()) == "THN"]
    if not direct_rows:
        return []
    candidates = conn.execute(sa_text("""
        SELECT emp_id, full_name, line_name
        FROM frontend.employee_groupings
        WHERE COALESCE(employee_status, 'ACTIVE') = 'ACTIVE'
    """)).mappings().all()
    # Kung ang agency ID mismo ay nasa upload na ito, sinadya ng Excel na
    # ilista siya -- huwag galawin.
    incoming_ids = {str(r["emp_id"]).strip().upper() for r in incoming_rows}
    stale = _find_transferred_stale(direct_rows, candidates, incoming_ids)
    _mark_inactive(conn, [s["emp_id"] for s in stale])
    for s in stale:
        print(f"[transfer-dedupe] {s['emp_id']} -> INACTIVE (napalitan ng {s['replaced_by']}: {s['full_name']})")
    return stale


def _upsert_batch_to_roster(conn, batch_id: int, all_employees: list) -> None:
    """Writes a batch's employee rows into the live frontend.employee_groupings
    roster. Called immediately for 'immediate' imports, or later by
    apply_due_imports() when a scheduled import's effective time arrives."""
    rows = _employee_rows(conn, batch_id, all_employees)
    if not rows:
        return
    # A late-applied scheduled batch or historical recovery upload must not
    # roll today's roster back over a newer effective assignment/manual move.
    newer = conn.execute(sa_text("""
        SELECT DISTINCT s.emp_id
        FROM frontend.employee_snapshot_records s
        JOIN frontend.import_batches b ON b.id = s.batch_id
        JOIN frontend.import_batches incoming ON incoming.id = :batch_id
        WHERE s.emp_id = ANY(:ids)
          AND b.effective_at <= (CURRENT_TIMESTAMP AT TIME ZONE 'Asia/Manila')
          AND (b.effective_at > incoming.effective_at
               OR (b.effective_at = incoming.effective_at AND b.id > incoming.id))
        UNION
        SELECT m.emp_id FROM frontend.employee_movements m
        JOIN frontend.import_batches incoming ON incoming.id = :batch_id
        WHERE m.emp_id = ANY(:ids) AND m.field_name = 'SHIFT'
          AND m.to_batch_id IS NULL AND m.detected_at >= incoming.effective_at
    """), {"batch_id": batch_id, "ids": [row["emp_id"] for row in rows]}).fetchall()
    newer_ids = {row[0] for row in newer}
    rows = [row for row in rows if row["emp_id"] not in newer_ids]
    if not rows:
        return
    conn.execute(sa_text("""
            INSERT INTO frontend.employee_groupings (
                emp_id, emp_no, emp, last_name, first_name, full_name,
                line_name, shift, process, line,
                department, position, machine, wing, remarks, gender,
                date_hired, tenure, updated_at
            ) VALUES (
                :emp_id, :emp_no, :emp, :last_name, :first_name, :full_name,
                :line_name, :shift, :process, :line,
                :department, :position, :machine, :wing, :remarks, :gender,
                :date_hired, :tenure, NOW()
            )
            ON CONFLICT (emp_id) DO UPDATE SET
                emp_no     = EXCLUDED.emp_no,
                emp        = EXCLUDED.emp,
                last_name  = EXCLUDED.last_name,
                first_name = EXCLUDED.first_name,
                full_name  = EXCLUDED.full_name,
                line_name  = EXCLUDED.line_name,
                shift      = EXCLUDED.shift,
                process    = EXCLUDED.process,
                line       = EXCLUDED.line,
                department = EXCLUDED.department,
                position   = EXCLUDED.position,
                machine    = EXCLUDED.machine,
                wing       = EXCLUDED.wing,
                remarks = EXCLUDED.remarks,
                gender = EXCLUDED.gender,
                date_hired = EXCLUDED.date_hired,
                tenure     = EXCLUDED.tenure,
                -- Kapag nasa Excel ulit ang dating INACTIVE (rehire / maling tanggal),
                -- ibalik sa ACTIVE. Ang ibang status ay hindi ginagalaw.
                employee_status = CASE
                    WHEN employee_groupings.employee_status = 'INACTIVE' THEN 'ACTIVE'
                    ELSE employee_groupings.employee_status
                END,
                updated_at = NOW()
        """), rows)
    deactivate_transferred_agency_rows(conn, rows)

# ── Auto-deactivate: mga wala na sa Excel ────────────────────────────────────
# Ang Excel na ina-upload ang source of truth (kumpletong file). Pagkatapos ng
# upsert, ang mga ACTIVE na nasa scope pero WALA sa file ay minamarkahang
# INACTIVE (hindi binubura -- buo pa rin ang attendance/movement history, at
# babalik sa ACTIVE kapag lumitaw ulit sa Excel).
#
# Mga safeguard:
#   * agency-scoped: "THN only" -> THN lang ang pwedeng ma-INACTIVE
#   * department-scoped: tinitingnan lang ang mga department na LAMAN ng file
#     (Department column, via parse_employee_department). Production file ->
#     Production lang; hindi madadamay ang Quality/Technical.
#   * authority-scoped: ang pwedeng ma-INACTIVE ay yung kayang makita ng
#     nag-upload (can_view_employee: department/section/line ng account niya).
#     Naka-save ang snapshot ng account sa batch para tama rin sa scheduled
#     imports, kung saan ibang user na ang tumatakbo ng apply.
#   * LINE-level accounts ay hindi pwedeng mag-deactivate (can_manage_employees_check)
#   * ang mga empleyado ng sheet na pumalya/walang laman/walang header
#     (protected_sheets) ay HINDI ginagalaw
#   * kapag walang nabasang empleyado sa file, walang ide-deactivate
#   * opt-in: kailangan ng deactivate_missing=1 mula sa uploader (may preview)
#   * kapag may mas bago nang batch na na-apply, hindi tatakbo (late scheduled)

_BATCH_EXTRA_COLUMNS = {
    "schedule_effective_at": "TIMESTAMP",
    "deactivate_scope":      "TEXT",
    "protected_sheets":      "TEXT",
    "deactivate_user":       "TEXT",
}


def ensure_batch_columns(conn) -> None:
    # ALTER TABLE ... ADD COLUMN IF NOT EXISTS takes an ACCESS EXCLUSIVE lock
    # even when the column exists, so it can stall behind any running query on
    # import_batches. Check the catalog first (no lock) and only ALTER if needed.
    have = {r[0] for r in conn.execute(sa_text("""
        SELECT column_name FROM information_schema.columns
        WHERE table_schema = 'frontend' AND table_name = 'import_batches'
          AND column_name = ANY(:cols)
    """), {"cols": list(_BATCH_EXTRA_COLUMNS)})}
    for col, ddl_type in _BATCH_EXTRA_COLUMNS.items():
        if col not in have:
            conn.execute(sa_text(
                f"ALTER TABLE frontend.import_batches ADD COLUMN IF NOT EXISTS {col} {ddl_type}"))


def _session_user_snapshot() -> dict:
    """Kopya ng access info ng nag-upload (para magamit ng can_view_employee
    kahit ibang user/session na ang tumatakbo, hal. scheduled apply)."""
    return {
        "role":             session.get("role"),
        "department":       session.get("department"),
        "access_level":     session.get("access_level"),
        "assigned_line":    session.get("assigned_line"),
        "assigned_shift":   session.get("assigned_shift"),
        "assigned_section": session.get("assigned_section"),
    }


def departments_in_import(all_employees: list) -> set:
    """Mga (parsed) department na laman ng Excel, hal. {"PRODUCTION"}."""
    return {parse_employee_department(e.get("department"))[0] for e in all_employees}


def _sheet_key(value: object) -> str:
    return str(value or "").strip().upper()


def find_employees_missing_from_import(conn, all_employees: list, scope: str,
                                       protected_sheets=(), user=None) -> list[dict]:
    """Mga ACTIVE na empleyado (sa scope) na wala sa import. Read-only.
    `user` = access snapshot ng nag-upload; kapag wala, walang ide-deactivate."""
    if user is None:
        return []
    present = {str(e.get("employee_id", "")).strip().upper() for e in all_employees}
    present.discard("")
    if not present:
        return []  # walang nabasa -> huwag mag-deactivate ng kahit sino

    file_depts = departments_in_import(all_employees)
    protected = {_sheet_key(s) for s in (protected_sheets or [])}
    rows = conn.execute(sa_text("""
        SELECT emp_id, full_name, line_name, shift, department
        FROM frontend.employee_groupings
        WHERE COALESCE(employee_status, 'ACTIVE') = 'ACTIVE'
        ORDER BY line_name, full_name
    """)).mappings().all()

    missing = []
    for r in rows:
        emp_id = str(r["emp_id"] or "").strip().upper()
        if not emp_id or emp_id in present:
            continue
        if scope != "ALL" and agency_code_from_emp_id(emp_id) != scope:
            continue
        if _sheet_key(r["line_name"]) in protected:
            continue
        if parse_employee_department(r["department"])[0] not in file_depts:
            continue  # ibang department -- hindi sakop ng file na ito
        if not can_view_employee({"department": r["department"],
                                  "line_name":  r["line_name"],
                                  "shift":      r["shift"]}, user=user):
            continue  # labas sa access ng nag-upload
        missing.append({
            "emp_id":    r["emp_id"],
            "full_name": r["full_name"] or "",
            "line_name": r["line_name"] or "",
            "shift":     r["shift"] or "",
        })
    return missing


def deactivate_missing_employees(conn, batch_id: int, all_employees: list) -> list[dict]:
    ensure_batch_columns(conn)
    row = conn.execute(sa_text("""
        SELECT deactivate_scope, protected_sheets, deactivate_user
        FROM frontend.import_batches WHERE id = :id
    """), {"id": batch_id}).fetchone()
    if not row or not row[0]:
        return []  # hindi hiniling ang deactivation (o lumang batch)
    if not row[2]:
        print(f"[auto-deactivate] batch {batch_id}: walang naka-save na uploader access, skipped")
        return []
    user = row[2] if isinstance(row[2], dict) else json.loads(row[2])

    scope = str(row[0]).strip().upper()
    raw_protected = row[1]
    if isinstance(raw_protected, list):
        protected = raw_protected
    else:
        protected = json.loads(raw_protected) if raw_protected else []

    newer = conn.execute(sa_text("""
        SELECT 1
        FROM frontend.import_batches b
        JOIN frontend.import_batches cur ON cur.id = :batch_id
        WHERE b.id <> cur.id AND b.applied = TRUE
          AND (b.effective_at > cur.effective_at
               OR (b.effective_at = cur.effective_at AND b.id > cur.id))
        LIMIT 1
    """), {"batch_id": batch_id}).fetchone()
    if newer:
        print(f"[auto-deactivate] batch {batch_id}: may mas bago nang batch, skipped")
        return []

    missing = find_employees_missing_from_import(conn, all_employees, scope, protected, user=user)
    _mark_inactive(conn, [m["emp_id"] for m in missing])
    for m in missing:
        print(f"[auto-deactivate] batch {batch_id}: {m['emp_id']} ({m['full_name']}, {m['line_name']}) -> INACTIVE")
    return missing


def apply_batch_to_roster(conn, batch_id: int, all_employees: list) -> list[dict]:
    """Upsert ng batch sa live roster, tapos auto-deactivate ng mga wala na
    sa Excel (kung hiniling). Ibinabalik ang listahan ng na-deactivate."""
    _upsert_batch_to_roster(conn, batch_id, all_employees)
    return deactivate_missing_employees(conn, batch_id, all_employees)


def apply_due_imports(conn) -> int:
    """Checks for scheduled import batches whose effective_at has passed
    and haven't been applied yet, and activates them: replays the roster
    upsert + movement detection + dated attendance requeue as an
    immediate import would have done. Called at the top of /api/data
    on every dashboard load. If a writer holds the lock, leave pending
    imports for a subsequent refresh and return the committed roster."""
    # Reads must not queue behind a full attendance-processing run.
    # Check cheaply first, then retain the shared lock for actual mutations.
    due = conn.execute(sa_text("""
        SELECT EXISTS (
            SELECT 1 FROM frontend.import_batches
            WHERE applied = FALSE
              AND effective_at <= (CURRENT_TIMESTAMP AT TIME ZONE 'Asia/Manila')
        )
    """)).scalar()
    if not due:
        return 0
    acquired = conn.execute(
        sa_text("SELECT pg_try_advisory_xact_lock(:key)"),
        {"key": ATTENDANCE_LOCK},
    ).scalar()
    if not acquired:
        # Leave the batch pending. The next dashboard refresh retries it;
        # meanwhile the reader can display the currently committed roster.
        return 0

    due_batches = conn.execute(sa_text("""
        SELECT id, raw_payload
        FROM frontend.import_batches
        WHERE applied = FALSE AND effective_at <= (CURRENT_TIMESTAMP AT TIME ZONE 'Asia/Manila')
        ORDER BY effective_at ASC, id ASC
    """)).fetchall()

    if not due_batches:
        return 0

    applied_count = 0
    for batch_id, raw_payload in due_batches:
        all_employees = raw_payload if isinstance(raw_payload, list) else json.loads(raw_payload)

        apply_batch_to_roster(conn, batch_id, all_employees)
        _, shift_changed = detect_movements(conn, batch_id, all_employees)
        queue_batch_attendance(conn, batch_id, all_employees)

        conn.execute(sa_text("""
            UPDATE frontend.import_batches SET applied = TRUE WHERE id = :id
        """), {"id": batch_id})

        applied_count += 1
        print(f"[apply_due_imports] Activated scheduled batch {batch_id}")

    return applied_count

@groupings_bp.route("/api/import-batches")
@login_required
def import_batches():
    # ── NEW: also return imported_by, effective_at, and applied so the
    # frontend can show "who uploaded what and when" without a second
    # round trip per batch.
    with db.engine.connect() as conn:
        rows = conn.execute(sa_text("""
            SELECT id, filename, imported_at, imported_by, effective_at, applied
            FROM frontend.import_batches
            ORDER BY imported_at DESC
        """)).mappings().all()
    return jsonify([{
        "id":           r["id"],
        "filename":     r["filename"],
        "imported_at":  str(r["imported_at"]),
        "imported_by":  r["imported_by"],
        "effective_at": str(r["effective_at"]) if r["effective_at"] else None,
        "applied":      r["applied"],
    } for r in rows])


# ── Content viewer: lets an Admin see what was actually in an uploaded Excel ──
# Columns shown in the "View" modal, in display order: (payload key, header label).
# Columns that are empty for every row in a sheet are dropped automatically.
BATCH_CONTENT_COLUMNS = [
    ("hc",          "HC"),
    ("emp_no",      "Emp No"),
    ("employee_id", "Employee ID"),
    ("full_name",   "Full Name"),
    ("last_name",   "Last Name"),
    ("first_name",  "First Name"),
    ("agency",      "Agency"),
    ("line",        "Line"),
    ("sub_line",    "Sub Line"),
    ("shift",       "Shift"),
    ("process",     "Process"),
    ("machine",     "Machine"),
    ("wing",        "Wing"),
    ("department",  "Department"),
    ("position",    "Position"),
    ("gender",      "Gender"),
    ("date_hired",  "Date Hired"),
    ("tenure",      "Tenure"),
    ("remarks",     "Remarks"),
]


def build_batch_content(payload) -> list[dict]:
    """Turns a batch's stored raw_payload (list of parsed employee dicts)
    into [{name, columns, rows}] grouped per Excel sheet, for the frontend."""
    employees = payload if isinstance(payload, list) else json.loads(payload or "[]")

    by_sheet: dict[str, list[dict]] = {}
    for emp in employees:
        by_sheet.setdefault(emp.get("sheet_name") or "Sheet", []).append(emp)

    sheets = []
    for sheet_name, emps in by_sheet.items():
        used = [
            (key, label) for key, label in BATCH_CONTENT_COLUMNS
            if any(e.get(key) not in (None, "") for e in emps)
        ]
        sheets.append({
            "name":    sheet_name,
            "columns": [label for _, label in used],
            "rows": [
                ["" if e.get(key) is None else e.get(key) for key, _ in used]
                for e in emps
            ],
        })
    return sheets


@groupings_bp.route("/api/import-batches/<int:batch_id>/content")
@admin_required
def import_batch_content(batch_id: int):
    """Admin-only: returns the employee rows that were imported in one batch,
    grouped per sheet. Read from the batch's stored raw_payload."""
    with db.engine.connect() as conn:
        row = conn.execute(sa_text("""
            SELECT filename, imported_by, imported_at, raw_payload
            FROM frontend.import_batches
            WHERE id = :id
        """), {"id": batch_id}).mappings().first()

    if not row:
        return jsonify({"success": False, "message": "Import batch not found."}), 404

    try:
        sheets = build_batch_content(row["raw_payload"])
    except (ValueError, TypeError):
        return jsonify({"success": False, "message": "Stored content for this batch is unreadable."}), 500

    return jsonify({
        "success":     True,
        "filename":    row["filename"],
        "imported_by": row["imported_by"],
        "imported_at": str(row["imported_at"]),
        "sheets":      sheets,
    })


@groupings_bp.route("/api/manual-move-employee", methods=["POST"])
@manage_employees_required
def manual_move_employee():
    try:
        data   = request.get_json(force=True, silent=True) or {}
        emp_id = str(data.get("emp_id", "")).strip().upper()

        if not emp_id:
            return jsonify({"success": False, "message": "Employee ID is required."}), 400

        with db.engine.begin() as conn:
            conn.execute(sa_text("SELECT pg_advisory_xact_lock(:key)"), {"key": ATTENDANCE_LOCK})
            movement_at = plant_now()
            existing = conn.execute(sa_text("""
                SELECT full_name, line, line_name, shift, process, department
                FROM frontend.employee_groupings WHERE emp_id = :emp_id
            """), {"emp_id": emp_id}).fetchone()

            if not existing:
                return jsonify({"success": False, "message": "Employee not found."}), 404

            full_name, old_line, old_line_name, old_shift, old_process, old_department = existing

            # Scope check: a Section/Department-level caller may only act on
            # employees within their own visibility scope, per
            # can_view_employee()'s docstring (every write route acting on
            # an employee_id must validate it server-side). Admin/GLOBAL
            # callers always pass. Return 404 rather than 403 so an
            # out-of-scope employee's existence isn't revealed.
            if not can_view_employee({
                "department": old_department,
                "line_name":  old_line_name,
                "shift":      old_shift,
            }):
                return jsonify({"success": False, "message": "Employee not found."}), 404

            final = {
                "line":      data.get("new_line")      or old_line,
                "line_name": data.get("new_line_name")  or old_line_name,
                "shift":     data.get("new_shift")     or old_shift,
                "process":   data.get("new_process")   or old_process,
            }

            conn.execute(sa_text("""
                UPDATE frontend.employee_groupings
                SET
                    line = :line,
                    line_name = :line_name,
                    shift = :shift,
                    process = :process,
                    updated_at = NOW()
                WHERE emp_id = :emp_id
            """), {**final, "emp_id": emp_id})

            field_map = {
                "LINE":      (old_line,      final["line"]),
                "LINE_NAME": (old_line_name, final["line_name"]),
                "SHIFT":     (old_shift,     final["shift"]),
                "PROCESS":   (old_process,   final["process"]),
            }
            for field_name, (old_val, new_val) in field_map.items():
                if old_val == new_val:
                    continue
                conn.execute(sa_text("""
                    INSERT INTO frontend.employee_movements
                        (emp_id, full_name, field_name, old_value, new_value, detected_at)
                    VALUES
                        (:emp_id, :full_name, :field_name, :old_value, :new_value, :movement_at)
                """), {
                    "emp_id": emp_id, "full_name": full_name,
                    "field_name": field_name, "old_value": old_val, "new_value": new_val,
                    "movement_at": movement_at,
                })

            # A manual line move can also change an employee's shift, so this
            # needs the same dated requeue as a bulk import gets. Runs on the
            # same connection/transaction as the update above.
            if old_shift != final["shift"]:
                queue_attendance_rebuild(conn, [emp_id], movement_at)

        return jsonify({"success": True, "employee": {"emp_id": emp_id, "full_name": full_name, **final}})

    except Exception as e:
        traceback.print_exc()
        return jsonify({"success": False, "message": str(e)}), 500

@groupings_bp.route("/api/all-lines")
@login_required
def all_lines():
    try:
        with db.engine.connect() as conn:
            rows = conn.execute(sa_text("""
                SELECT DISTINCT line_name, line, shift
                FROM frontend.employee_groupings
                WHERE line IS NOT NULL AND line != ''
                ORDER BY line_name, line
            """)).mappings().all()

        return jsonify([{
            "line_name": r["line_name"], "line": r["line"], "shift": r["shift"],
        } for r in rows])

    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@groupings_bp.route("/api/employee-history/<emp_id>")
@login_required
def employee_history(emp_id: str):
    try:
        with db.engine.connect() as conn:
            rows = conn.execute(sa_text("""
                SELECT
                    m.field_name, m.old_value, m.new_value,
                    b_from.filename AS from_file, b_from.imported_at AS from_date,
                    b_to.filename   AS to_file,   b_to.imported_at   AS to_date,
                    m.detected_at
                FROM frontend.employee_movements m
                LEFT JOIN frontend.import_batches b_from ON b_from.id = m.from_batch_id
                LEFT JOIN frontend.import_batches b_to   ON b_to.id   = m.to_batch_id
                WHERE m.emp_id = :emp_id
                ORDER BY m.detected_at DESC
            """), {"emp_id": emp_id}).fetchall()

        return jsonify([{
            "field":       r[0], "from_value": r[1], "to_value":   r[2],
            "from_file":   r[3], "from_date":  str(r[4]) if r[4] else None,
            "to_file":     r[5], "to_date":    str(r[6]) if r[6] else None,
            "detected_at": str(r[7]),
        } for r in rows])

    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@groupings_bp.route("/api/movements-summary")
@login_required
def movements_summary():
    try:
        with db.engine.connect() as conn:
            latest = conn.execute(sa_text("""
                SELECT id, filename, imported_at
                FROM frontend.import_batches
                ORDER BY id DESC LIMIT 1
            """)).fetchone()

            if not latest:
                return jsonify({"movements": [], "batch": None})

            rows = conn.execute(sa_text("""
                SELECT m.emp_id, m.full_name, m.field_name,
                       m.old_value, m.new_value, m.detected_at
                FROM frontend.employee_movements m
                WHERE m.to_batch_id = :bid
                ORDER BY m.full_name
            """), {"bid": latest[0]}).fetchall()

        return jsonify({
            "batch": {"id": latest[0], "filename": latest[1], "imported_at": str(latest[2])},
            "movements": [{
                "emp_id": r[0], "full_name": r[1], "field": r[2],
                "from_value": r[3], "to_value": r[4], "detected_at": str(r[5]),
            } for r in rows],
        })

    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@groupings_bp.route("/api/admin/cleanup-transferred-duplicates", methods=["POST"])
@admin_required
def cleanup_transferred_duplicates():
    """Isang beses na cleanup ng mga doble na nasa DB na (agency -> THN).
    Default ay DRY RUN (walang binabago, listahan lang ang ibabalik).
    Para ituloy: POST /api/admin/cleanup-transferred-duplicates?dry_run=0"""
    dry_run = str(request.args.get("dry_run", "1")).strip().lower() not in ("0", "false", "no")
    with db.engine.begin() as conn:
        rows = conn.execute(sa_text("""
            SELECT emp_id, full_name, line_name
            FROM frontend.employee_groupings
            WHERE COALESCE(employee_status, 'ACTIVE') = 'ACTIVE'
        """)).mappings().all()
        direct = [r for r in rows if agency_code_from_emp_id(str(r["emp_id"]).strip().upper()) == "THN"]
        stale = _find_transferred_stale(direct, rows)
        if not dry_run:
            _mark_inactive(conn, [s["emp_id"] for s in stale])
    return jsonify({"success": True, "dry_run": dry_run, "count": len(stale), "rows": stale})